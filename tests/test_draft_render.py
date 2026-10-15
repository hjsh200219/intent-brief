import os
import re
import zipfile

import pytest

import _paths as P
import check_draft as CD
import intent as I
import render_docx as R
import state as S
from conftest import F1, F1_BASE, make_state

ADMIT2 = [{"id": 1, "title": "금전소비대차의 성립", "stance": "다툼", "evidence": {"mode": "제출", "items": ["갑1", "갑2"]}},
          {"id": 2, "title": "변제 항변", "stance": "일부 인정", "amount": 10000000, "evidence": {"mode": "유보"}}]


def design_s_m(params):
    st = S.new_state(F1, "준비서면", params)
    st.update(side="원고", claim_amount=50000000, issues=[dict(i) for i in ADMIT2])
    st["intent"].update(primary="S", secondaries=["M"], success_sentence="s", concession="4,000만원")
    st["flags"]["settle_will"] = True
    return I.resolve(st, params)


def good_draft(design, filler=180):
    body = "원고는 피고와의 원만한 해결을 바랍니다. " * filler
    parts = ["---", "design: x", "side: 원고", "brief_type: 준비서면", "---",
             "## [sec:sec.header] 준비서면", "사건 2026가단512345 대여금", "원고 홍길동", "피고 김대출"]
    for s in design["sections"][1:]:
        parts.append("## [sec:%s] %s" % (s["id"], s["title"]))
        if s["id"] == "sec.background":
            parts.append(body)
        elif s["id"] == "sec.admitted_facts":
            parts.append("원고는 피고가 1,000만원을 변제한 사실을 인정합니다.")
        elif s["id"] == "sec.rebuttal_core":
            parts += ["### [issue:1] 금전소비대차의 성립", "피고의 증여 주장은 사실이 아닙니다. 원고는 이를 부인합니다.",
                      "### [issue:2] 변제 항변", "원고는 1,000만원 변제 사실을 인정합니다."]
        elif s["id"] == "sec.conclusion":
            parts.append(design["conclusion"]["text"])
        elif s["id"] == "sec.evidence_list":
            parts += ["1. 갑 제1호증 계좌이체내역 [확인 필요: 이체일]", "0000. 00. 00.", "원고 소송대리인 변호사 ○○○ (인)"]
        else:
            parts.append("내용")
    return "\n".join(parts) + "\n"


def test_good_draft_full_compliance(params):
    d = design_s_m(params)
    res = CD.check(good_draft(d), d)
    assert res["failures"] == [] and res["rate"] == 1.0, res["failures"]
    assert res["metrics"]["check_tags"] == 1


def test_detects_violations(params):
    d = design_s_m(params)
    md = good_draft(d)
    md = md.replace("## [sec:sec.settle_opinion]", "## 조정 의견")  # 앵커 없음 + 섹션 누락
    md = md.replace("원고는 1,000만원 변제 사실을 인정합니다.", "피고의 변제 주장은 근거가 없습니다.")
    md = md.replace(d["conclusion"]["text"], "원고의 청구를 전부 인용하여 주시기 바랍니다")
    md += "## [sec:sec.rebuttal_all] 개별 반박\n명백히 전혀 결코 틀림없이 허위입니다.\n"
    ids = {f["id"] for f in CD.check(md, d)["failures"]}
    assert {"anchors", "section:sec.settle_opinion", "stance:2", "conclusion", "no_extra_sections", "strength"} <= ids


def test_reserved_issue_must_not_appear(params):
    d = design_s_m(params)
    d["issues"][1]["stance"] = "유보"
    ids = {f["id"] for f in CD.check(good_draft(d), d)["failures"]}
    assert "stance:2" in ids


def test_amount_required_for_partial(params):
    d = design_s_m(params)
    md = good_draft(d).replace("원고는 1,000만원 변제 사실을 인정합니다.", "원고는 변제 사실 일부를 인정합니다.")
    f = {x["id"]: x for x in CD.check(md, d)["failures"]}
    assert "stance:2" in f and "1,000만원" in f["stance:2"]["detail"]


def test_length_bounds(params):
    d = design_s_m(params)
    ids = {f["id"] for f in CD.check(good_draft(d, filler=2), d)["failures"]}
    assert "length" in ids


def test_record_first_final(home, case, params):
    c = case(F1)
    path, _ = make_state(c, **F1_BASE)
    d = design_s_m(params)
    wd = os.path.dirname(path)
    P.atomic_write_json(os.path.join(wd, "d.json"), d)
    P.atomic_write_text(os.path.join(wd, "draft.md"), good_draft(d))
    CD.main(["check", os.path.join(wd, "draft.md"), os.path.join(wd, "d.json"), "--state", path, "--record", "first"])
    st = P.load_json(path)
    assert st["compliance"]["first"]["rate"] == 1.0 and st["draft_path"].endswith("draft.md")


# ── docx ──────────────────────────────────────────────────────
@pytest.fixture
def rendered(home, case, params):
    c = case(F1)
    path, _ = make_state(c, **F1_BASE)
    st = P.load_json(path)
    st["intent"].update(primary="S", secondaries=["M", "C"], success_sentence="s", concession="4,000만원")
    st["issues"] = [dict(i) for i in ADMIT2]
    st["flags"]["settle_will"] = True
    st["intent_review"] = {"for": P.intent_signature(st["intent"]), "verdict": "조건부", "gains": ["조기 종결"],
                           "forgoes": ["항소심 청구취지 확장"], "alternatives": ["확장 신청 후 조정"], "decision": "진행",
                           "note": "의뢰인이 신속 종결 원함"}
    d = I.resolve(st, params)
    wd = os.path.dirname(path)
    st["design_path"] = os.path.join(wd, "준비서면_design.json")
    st["draft_path"] = os.path.join(wd, "준비서면_draft.md")
    st["status"] = "designed"
    P.atomic_write_json(st["design_path"], d)
    P.atomic_write_text(st["draft_path"], good_draft(d) + "\n| 상대 주장 | 원고 의견 |\n|---|---|\n| 증여 | 부인 |\n")
    P.atomic_write_text(os.path.join(wd, "client_report.md"), "이번 서면은 조정을 목표로 합니다.\n")
    S.save(path, st)
    out = R.render_all(path, "대여금")
    return c, path, out


def test_render_files_and_names(rendered):
    c, path, out = rendered
    names = sorted(os.listdir(c))
    assert "260915_대여금_준비서면_v1.docx" in names
    assert "내부_260915_대여금_준비서면_v1_의도대응표.docx" in names
    assert "내부_260915_대여금_준비서면_v1_의뢰인보고.docx" in names
    assert P.load_json(path)["status"] == "drafted"
    out2 = R.render_all(path, "대여금")
    assert out2["version"] == 2  # 덮어쓰지 않고 vN 증가
    assert os.path.exists(os.path.join(c, "260915_대여금_준비서면_v1.docx"))


def test_docx_format_python_docx(rendered):
    import docx
    from docx.shared import Mm
    c, path, out = rendered
    doc = docx.Document(out["paths"]["draft"])
    sec = doc.sections[0]
    assert abs(sec.page_width - Mm(210)) < Mm(1) and abs(sec.page_height - Mm(297)) < Mm(1)
    assert abs(sec.top_margin - Mm(25)) < Mm(0.5) and abs(sec.right_margin - Mm(20)) < Mm(0.5)
    assert abs(sec.bottom_margin - Mm(25)) < Mm(0.5) and abs(sec.left_margin - Mm(25)) < Mm(0.5)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "[sec:" not in text and "[issue:" not in text  # 앵커 제거
    assert "인정한 1,000만원을 공제한 나머지 4,000만원을 인용" in text
    reds = [r.text for p in doc.paragraphs for r in p.runs if r.font.color is not None and r.font.color.rgb is not None and str(r.font.color.rgb) == "C00000"]
    assert reds == ["[확인 필요: 이체일]"]
    assert len(doc.tables) == 1 and doc.tables[0].cell(1, 1).text == "부인"
    t = docx.Document(out["paths"]["table"])
    alltext = "\n".join(p.text for p in t.paragraphs) + "\n".join(c.text for tb in t.tables for row in tb.rows for c in row.cells)
    assert "내부용" in alltext and "S" in alltext and "2016다225353" in alltext and "적합성" in alltext
    assert "의도 유불리 검토" in alltext and "조건부" in alltext and "항소심 청구취지 확장" in alltext and "진행" in alltext


def test_ooxml_structure(rendered):
    c, path, out = rendered
    z = zipfile.ZipFile(out["paths"]["draft"])
    names = set(z.namelist())
    assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml", "word/settings.xml",
            "word/_rels/document.xml.rels", "docProps/core.xml", "docProps/app.xml"} <= names
    ct = z.read("[Content_Types].xml").decode()
    for part in ("/word/document.xml", "/word/styles.xml", "/word/settings.xml", "/docProps/core.xml", "/docProps/app.xml"):
        assert 'PartName="%s"' % part in ct
    styles = z.read("word/styles.xml").decode("utf-8")
    assert 'w:eastAsia="바탕"' in styles and '<w:sz w:val="24"/>' in styles and '<w:jc w:val="both"/>' in styles
    doc = z.read("word/document.xml").decode("utf-8")
    # 스키마 순서: pPr 안 pStyle 이 jc 보다 앞, rPr 안 b 가 color 보다 앞, sectPr 은 body 마지막
    for m in re.finditer(r"<w:pPr>(.*?)</w:pPr>", doc):
        inner = m.group(1)
        if "<w:pStyle" in inner and "<w:jc" in inner:
            assert inner.index("<w:pStyle") < inner.index("<w:jc")
        if "<w:spacing" in inner and "<w:jc" in inner:
            assert inner.index("<w:spacing") < inner.index("<w:jc")
    for m in re.finditer(r"<w:rPr>(.*?)</w:rPr>", doc):
        inner = m.group(1)
        if "<w:b/>" in inner and "<w:color" in inner:
            assert inner.index("<w:b/>") < inner.index("<w:color")
    assert doc.rstrip().endswith("</w:sectPr></w:body></w:document>")
    for m in re.finditer(r"<w:tc>(.*?)</w:tc>", doc):
        assert m.group(1).startswith("<w:tcPr>") and "<w:p" in m.group(1)


def test_no_overwrite(tmp_path):
    p = str(tmp_path / "a.docx")
    R.write_docx(p, R.para("x"))
    with pytest.raises(P.BriefError):
        R.write_docx(p, R.para("y"))


def test_render_needs_design_and_draft(home, case):
    path, _ = make_state(case(F1))
    with pytest.raises(P.BriefError):
        R.render_all(path, "x")
